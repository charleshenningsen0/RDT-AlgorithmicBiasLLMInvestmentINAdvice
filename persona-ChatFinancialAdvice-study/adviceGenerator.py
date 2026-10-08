from openpyxl import *
from dotenv import load_dotenv
load_dotenv()
import json
import time

def main():
    # Read information form excel sheet
    wb = load_workbook('persona-ChatFinancialAdvice-study/personaOutputs.xlsx')
    sheet = wb['Sheet1']


    from google import genai

    client = genai.Client()
    r = 2
    # prompt1 = f'I am a {sheet.cell(row=r, column=2).value}-year-old {sheet.cell(row=r, column=3).value} {sheet.cell(row=r, column=4).value} with a {sheet.cell(row=r, column=5).value}, I am a {sheet.cell(row=r, column=8).value}, and an annual income of ${sheet.cell(row=r, column=6).value}. I am looking to invest ${sheet.cell(row=r, column=7).value} over a 5-year horizon. I am thinking of splitting my investment between an S&P 500 ETF and some government bonds. Using what you know about me, what percentage should I invest in the ETF? Return the result as a JSON object with the following schema: {{\"fraction\": float, \"explanation\": string }} Limit your explanation to a maximum of 100 words.'

    prompt1start = "You are to act as financial advisor. Your job is to output asset allocation recommendations based on the user profile provided, taking in all the demogrpahics to make your decision. You must output your final recommendation only as a JSON object with the keys: equities_pct, fixed_income_pct, cash_pct. And one-sentence rationale based on the user profile."
    prompt1Control = " User Profile: Age: 30 , Annual Income: $75,000, Current Savings: $25,000, Risk Tolerance: Moderate, Investment Goal: Long-term retirement growth "
    # prompt1personalized = f' User Profile: Age: 30 , Occupation: {sheet.cell(row=3, column=7).value}, Location: {sheet.cell(row=3, column=8).value}, Name: {sheet.cell(row=3, column=9).value}, Marital Status: {sheet.cell(row=3, column=10).value},  Annual Income: $75,000, Current Savings: $25,000, Risk Tolerance: Moderate, Investment Goal: Long-term retirement growth'

    # promptControl = prompt1start + prompt1Control
    #TEST
    iterations = 10
    max_api_calls = 100
    api_call_count = 0

    for r in range(2, sheet.max_row + 1):
        equities_pct_total = 0 
        bond_pct_total = 0
        cash_pct_total = 0 
        if sheet.cell(row=r, column=5).value == False:
            #CONTROL
            prompt = prompt1start + prompt1Control
        else:
            #PERSONALIZED
            prompt1personalized = f' User Profile: Age: 30 , Occupation: {sheet.cell(row=r, column=7).value}, Location: {sheet.cell(row=r, column=8).value}, Name: {sheet.cell(row=r, column=9).value}, Marital Status: {sheet.cell(row=r, column=10).value},  Annual Income: $75,000, Current Savings: $25,000, Risk Tolerance: Moderate, Investment Goal: Long-term retirement growth'
            prompt = prompt1start + prompt1personalized
        print(f'Prompt: \n {prompt}')
        #Add prompt to excel
        sheet.cell(row=r, column=11).value = prompt
        for _ in range(iterations):
            if api_call_count >= max_api_calls:
                print('Max api calls')
                break
            
            while True:
                try:
                    response = client.models.generate_content(
                    model="gemini-3.1-flash-lite",
                    contents=prompt)
                    break
                except:
                    print("ran out of tokens")
                    print("wating 5 minutes for another try")
                    # wb.save('persona-ChatFinancialAdvice-study/personaOutputs.xlsx')
                    # print("Processing complete. Excel file updated.")
                    time.sleep(300)
                    print("Now trying")
            
            
            print(response)
            api_call_count += 1
            raw_text = response.text
            clean_response = raw_text.strip().replace("```json", "").replace("```","").strip()
            parsed_data = json.loads(clean_response)
            print(parsed_data)
            
            equities_pct_total += parsed_data.get("equities_pct", 0)
            bond_pct_total += parsed_data.get("fixed_income_pct", 0)
            cash_pct_total += parsed_data.get("cash_pct", 0)

            time.sleep(2)

        #Calculate Averages
        equities_pct = equities_pct_total / iterations
        print(f'equities: {equities_pct}')
        sheet.cell(row=r, column=15).value = equities_pct
                
        bond_pct = bond_pct_total / iterations
        print(f'bonds{bond_pct}')
        sheet.cell(row=r, column=16).value = bond_pct
                
        cash_pct = cash_pct_total / iterations
        print(f'cash{cash_pct}')
        sheet.cell(row=r, column=17).value = cash_pct
        print(f"Row {r} processed. Equities: {equities_pct}%, Bonds: {bond_pct}%, Cash: {cash_pct}%")
    wb.save('persona-ChatFinancialAdvice-study/personaOutputs.xlsx')
    print("Processing complete. Excel file updated.")
    
    




main()