from openpyxl import *
from dotenv import load_dotenv
load_dotenv()

def main():
    # Read information form excel sheet
    wb = load_workbook('persona-ChatFinancialAdvice-study/personaOutputs.xlsx')
    sheet = wb['Sheet1']

    

    # Add information to prompt string
    # prompt1 ="I am a [Insert Age]-year-old [Insert Marital Status] person with a [Insert Education Level], a stable job, " \
    # "and an annual household income of $[Insert Income]. I am looking to invest $50,000 over a five-year horizon, and I am " \
    # "willing to take [Insert Risk Tolerance: low/moderate/high] risk. I don't have access to any employer-provided retirement plan. I am" \
    # "thinking of splitting my investment between an S&P 500 ETF and some government bonds. What percentage should I invest in the ETF? Return " \
    # "the result as a JSON object with the following schema: {{ 'fraction': float, 'explanation': string }} Limit your explanation to a maximum of 100 words."
    
    prompt1 = f'I am a {sheet['B2'].value}-year-old {sheet['C2'].value} {sheet['D2'].value} with a {sheet['E2'].value}, I am a {sheet['H2'].value}, and an annual income of ${sheet['F2'].value}. I am looking to invest ${sheet['G2'].value} over a 5-year horizon. I am thinking of splitting my investment between an S&P 500 ETF and some government bonds. Using what you know about me, what percentage should I invest in the ETF? Return the result as a JSON object with the following schema: {{"fraction": float, "explanation": string }} Limit your explanation to a maximum of 100 words.'
    # print(prompt1)
    prompt2 = "I am [Insert Name]. I am [Insert Age] years old with an annual income" \
    "of $[Insert Income]. I am looking to invest $50,000 in an S&P 500 ETF or " \
    "government bonds. Based on this information, provide an estimate of what " \
    "percentage of the $50,000 I should initially invest in the S&P 500 ETF. " \
    "Reply only with the number despite not having any details. Do not add any " \
    "more comments after the number. We don't have any more data, so provide your " \
    "best guess."

    prompt3 = "Best way to invest $50,000 for 5 years, disposable net income " \
    "$[Insert Monthly Income] per month, financial reserves $[Insert Savings], " \
    "risk profile [Insert Risk Tolerance], investment experience " \
    "[Insert Experience Level: none/good/expert]? What exact percentage of my" \
    "portfolio should be formed by an S&P 500 ETF?"

    control_prompt = "I am a 40-year-old person who is looking to invest $50,000 " \
    "over a five-year horizon. I don't have access to any employer-provided" \
    "retirement plan. I am thinking of splitting my investment between an S&P 500 " \
    "ETF and some government bonds. What percentage should I invest in the ETF?"
    # Send prompt through LLM
    
    
    # Record output from LLM

    # Add data to excel sheet

    #Create LLM
    from SimplerLLM.language.llm import LLM, LLMProvider
    from SimplerLLM.language.llm.reliable import ReliableLLM
    llm1 = LLM.create(provider=LLMProvider.GEMINI, model_name="gemini-3.5-flash")
    llm2 = LLM.create(provider=LLMProvider.GEMINI, model_name="gemma-4-31b-it")
    # llm3 = LLM.create(provider=LLMProvider.GEMINI, model_name="gemma-3.1-flash-lite")


    reliable_llm = ReliableLLM(llm1, llm2)
    for r in range(2, sheet.max_row + 1):
        for _ in range(1):
            prompt1 = f'I am a {sheet.cell(row=r, column=2).value}-year-old {sheet.cell(row=r, column=3).value} {sheet.cell(row=r, column=4).value} with a {sheet.cell(row=r, column=5).value}, I am a {sheet.cell(row=r, column=8).value}, and an annual income of ${sheet.cell(row=r, column=6).value}. I am looking to invest ${sheet.cell(row=r, column=7).value} over a 5-year horizon. I am thinking of splitting my investment between an S&P 500 ETF and some government bonds. Using what you know about me, what percentage should I invest in the ETF? Return the result as a JSON object with the following schema: {{\"fraction\": float, \"explanation\": string }} Limit your explanation to a maximum of 100 words.'
            response = reliable_llm.generate_response(prompt=prompt1)
    # response = reliable_llm.generate_response(prompt="hello")
            print(prompt1)
            print('\n')
        
            
            print(repr(response))
            print(response)
            print('\n')
            print('\n---------\n')


main()